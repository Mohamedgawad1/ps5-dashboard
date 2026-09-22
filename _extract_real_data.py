import openpyxl
import json
import os
import datetime

EXCEL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PS-5 COMPLETIONS DPR SUMMERY - 19-08-26.xlsx")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dashboard_data.json")

wb = openpyxl.load_workbook(EXCEL, data_only=True)

data = {"date": "19 AUGUST 2026", "disciplines": {}, "punch_by_cat": {}, "itr_summary": {}}

# ===== ITR SUMMARY =====
ws = wb["ITR SUMMARY"]
cpp_t = [0,0,0]
eit_t = [0,0,0]
disc_map = {"Building":"B","Electrical":"E","HVAC":"H","Instrumentation":"I","Mechanical":"M","Piping & Vessel":"P","Safety":"S","Telecom":"T"}

for row in ws.iter_rows(min_row=9, max_row=13, values_only=True):
    name, total, closed, pct, opn = row[0], row[1], row[2], row[3], row[4]
    if name and total:
        code = disc_map.get(name, name[0])
        data["disciplines"][code] = {"total": total, "closed": closed, "in_progress": 0, "pending": opn}
        cpp_t[0] += total; cpp_t[1] += closed; cpp_t[2] += opn

for row in ws.iter_rows(min_row=17, max_row=19, values_only=True):
    name, total, closed, pct, opn = row[0], row[1], row[2], row[3], row[4]
    if name and total:
        code = disc_map.get(name, name[0])
        data["disciplines"][code] = {"total": total, "closed": closed, "in_progress": 0, "pending": opn}
        eit_t[0] += total; eit_t[1] += closed; eit_t[2] += opn

grand = [cpp_t[0]+eit_t[0], cpp_t[1]+eit_t[1], cpp_t[2]+eit_t[2]]
data["itr_summary"] = {"cpp": cpp_t, "eit": eit_t, "grand": grand}

# ===== PUNCH SUMMARY =====
ws2 = wb["PUNCH SUMMERY"]
disc_order = ["Building","Electrical","HVAC","Instrumentation","Mechanical","Piping & Vessel","Safety","Telecom"]
# Columns: A=3, B=9, C=15, 0=21 (each has Discipline, Total, Closed, Open)
cat_map = {3: "A", 9: "B", 15: "C", 21: "0"}

for cat_col, cat_letter in cat_map.items():
    for d_idx, disc_name in enumerate(disc_order):
        r = 7 + d_idx
        total_cell = ws2.cell(row=r, column=cat_col+1).value or 0
        closed_cell = ws2.cell(row=r, column=cat_col+2).value or 0
        open_cell = ws2.cell(row=r, column=cat_col+3).value or 0
        
        code = disc_map.get(disc_name, disc_name[0])
        if code not in data["punch_by_cat"]:
            data["punch_by_cat"][code] = {"total": 0, "closed": 0, "open": 0}
        data["punch_by_cat"][code]["total"] += total_cell
        data["punch_by_cat"][code]["closed"] += closed_cell
        data["punch_by_cat"][code]["open"] += open_cell

# ===== SUBSYSTEMS from RFC PROGRESS =====
ws3 = wb["RFC PROGRESS"]
subs = []
for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, values_only=True):
    if row[0]:
        parts = str(row[0]).split(" - ", 1)
        sid = parts[0].strip()
        desc = parts[1].strip() if len(parts) > 1 else ""
        subs.append({"id": sid, "d": desc, "full": str(row[0])})
data["subsystems"] = subs
data["subsystems_count"] = len(subs)

# ===== MILESTONE DATA =====
ws4 = wb["MILESTONE_12 AUG 2026"]
milestone_data = []
for row in ws4.iter_rows(min_row=2, max_row=ws4.max_row, values_only=True):
    if row[0]:
        milestone_data.append({
            "id": str(row[0]),
            "desc": str(row[1]) if row[1] else "",
            "full": str(row[2]) if row[2] else "",
            "priority": str(row[3]) if row[3] else "",
            "month": str(row[4]) if row[4] else "",
        })
data["milestones"] = milestone_data

# ===== PUNCH SEVERITY from PUNCH SUMMARY =====
ws5 = wb["PUNCH SUMMERY"]
punch_severity = {}
for cat_col, cat_letter in cat_map.items():
    punch_severity[cat_letter] = {"total": 0, "closed": 0, "open": 0}
    for d_idx, disc_name in enumerate(disc_order):
        r = 7 + d_idx
        t = ws5.cell(row=r, column=cat_col+1).value or 0
        c = ws5.cell(row=r, column=cat_col+2).value or 0
        o = ws5.cell(row=r, column=cat_col+3).value or 0
        punch_severity[cat_letter]["total"] += t
        punch_severity[cat_letter]["closed"] += c
        punch_severity[cat_letter]["open"] += o
data["punch_severity"] = punch_severity

# ===== MILESTONE & DIS. =====
ws6 = wb["MILESTONE & DIS."]
milestone_disc = {}
for row in ws6.iter_rows(min_row=8, max_row=15, values_only=True):
    name, total, closed, pct, opn = row[0], row[1], row[2], row[3], row[4]
    if name and total:
        code = disc_map.get(name, name[0])
        milestone_disc[code] = {"total": total, "closed": closed, "pct": pct, "open": opn}
data["milestone_disc"] = milestone_disc

# ===== RFC PROGRESS per subsystem =====
ws7 = wb["RFC PROGRESS"]
rfc_data = []
for row in ws7.iter_rows(min_row=2, max_row=ws7.max_row, values_only=True):
    if row[0]:
        rfc_data.append({
            "subsystem": str(row[0]),
            "priority": row[1],
            "rfc_date": str(row[2]) if row[2] else "",
            "milestone": str(row[3]) if row[3] else "",
            "total_pct": row[4],
            "b_total": row[5], "b_closed": row[6], "b_bal": row[7],
            "h_total": row[8], "h_closed": row[9], "h_bal": row[10],
            "m_total": row[11], "m_closed": row[12], "m_bal": row[13],
            "p_total": row[14], "p_closed": row[15], "p_bal": row[16],
            "s_total": row[17], "s_closed": row[18], "s_bal": row[19],
            "total_itrs": row[20], "total_closed": row[21], "total_bal": row[22],
            "itr_pct": row[23],
            "remarks": str(row[24]) if row[24] else ""
        })
data["rfc_progress"] = rfc_data

# ===== EIT PROGRESS per subsystem =====
ws8 = wb["EIT PROGRESS"]
eit_data = []
for row in ws8.iter_rows(min_row=2, max_row=ws8.max_row, values_only=True):
    if row[0]:
        eit_data.append({
            "subsystem": str(row[0]),
            "priority": row[1],
            "rfc_date": str(row[2]) if row[2] else "",
            "milestone": str(row[3]) if row[3] else "",
            "e_total": row[4], "e_closed": row[5], "e_remain": row[6], "e_pct": row[7],
            "i_total": row[8], "i_closed": row[9], "i_remain": row[10], "i_pct": row[11],
            "t_total": row[12], "t_closed": row[13], "t_remain": row[14], "t_pct": row[15],
            "total": row[16], "closed": row[17], "balance": row[18], "pct": row[19],
            "remarks": str(row[22]) if row[22] else ""
        })
data["eit_progress"] = eit_data

# ===== M-P PROGRESS per subsystem =====
ws9 = wb["M-P PROGRESS"]
mp_data = []
for row in ws9.iter_rows(min_row=2, max_row=ws9.max_row, values_only=True):
    if row[0]:
        mp_data.append({
            "subsystem": str(row[0]),
            "priority": row[1],
            "rfc_date": str(row[2]) if row[2] else "",
            "milestone": str(row[3]) if row[3] else "",
            "b_total": row[4], "b_closed": row[5], "b_bal": row[6], "b_pct": row[7],
            "h_total": row[8], "h_closed": row[9], "h_bal": row[10], "h_pct": row[11],
            "m_total": row[12], "m_closed": row[13], "m_bal": row[14], "m_pct": row[15],
            "p_total": row[16], "p_closed": row[17], "p_bal": row[18], "p_pct": row[19],
            "s_total": row[20], "s_closed": row[21], "s_bal": row[22], "s_pct": row[23],
            "total": row[24], "closed": row[25], "balance": row[26], "pct": row[27],
            "remarks": str(row[32]) if row[32] else ""
        })
data["mp_progress"] = mp_data

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(data, f, indent=2, default=str)

print(f"Real data extracted to {OUT}")
print(f"Subsystems: {len(subs)}")
print(f"RFC rows: {len(rfc_data)}")
print(f"EIT rows: {len(eit_data)}")
print(f"MP rows: {len(mp_data)}")
print(f"Itr summary: grand={grand}")
print(f"Punch by discipline: {json.dumps(data['punch_by_cat'], indent=2)}")
