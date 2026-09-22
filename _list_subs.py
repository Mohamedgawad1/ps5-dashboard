import openpyxl, json

wb = openpyxl.load_workbook(r'C:\Users\mylap\OneDrive\Desktop\dashboard\PS-5 COMPLETIONS DPR SUMMERY - 19-08-26.xlsx', data_only=True)

# Get all 66 subsystems with descriptions
ws = wb['DETAILED PUNCH LIST']
sub_map = {}
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    sid = row[7]  # Systemization - Subsystem (Summary)
    if sid and sid not in sub_map:
        parts = str(sid).split(' - ', 1)
        sub_id = parts[0].strip()
        desc = parts[1].strip() if len(parts) > 1 else ''
        sub_map[sid] = {'id': sub_id, 'd': desc, 'full': str(sid)}

# Print all
for i, (k, v) in enumerate(sorted(sub_map.items())):
    print(f"{i+1}. {v['id']} | {v['d']}")

print(f"\nTotal: {len(sub_map)}")
