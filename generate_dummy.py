import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import random
import datetime
import json
import os

random.seed(42)

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
TODAY = datetime.date.today()
TODAY_STR = TODAY.strftime("%d %B %Y").upper()
DATE_FILE = TODAY.strftime("%d-%m-%y")  # e.g. 19-08-26

# ===== BEIGE / GOLDEN COLOR SCHEME =====
C_GOLD_DARK    = "8B6914"
C_GOLD_MED     = "C8A96E"
C_GOLD_LIGHT   = "E8D5A3"
C_CREAM        = "FDF6E3"
C_BEIGE_BG     = "FAF3E0"
C_BROWN_DARK   = "4A3728"
C_BROWN_MED    = "6B5344"
C_BROWN_LIGHT  = "8B7355"
C_WHITE        = "FFFFFF"
C_ROW_1        = "FDF8EF"
C_ROW_2        = "FAF3E0"
C_ACCENT_GREEN = "6B8E23"
C_ACCENT_RED   = "B22222"
C_ACCENT_BLUE  = "4682B4"

DISCIPLINES = {
    "B": {"name": "Building",        "color": "C8A96E"},
    "E": {"name": "Electrical",      "color": "D4A847"},
    "H": {"name": "HVAC",            "color": "A08050"},
    "I": {"name": "Instrumentation", "color": "DAA520"},
    "M": {"name": "Mechanical",      "color": "B8860B"},
    "P": {"name": "Piping & Vessel", "color": "8B7355"},
    "S": {"name": "Safety",          "color": "CD853F"},
    "T": {"name": "Telecom",         "color": "C49E6C"},
}

COMPANIES = ["CPP AGI", "EPS UNATRAC", "KENT", "EITS", "MANTRAC"]
MILESTONES = ["PS5 - Milestone A", "PS5 - Milestone B", "PS5 - Milestone C",
              "PS5 - Milestone D", "PS5 - Milestone E", "PS5 - Milestone F",
              "PS5 - Milestone G", "PS5 - Milestone H", "PS5 - Milestone I"]
MONTHS = ["JANUARY","FEBRUARY","MARCH","APRIL","MAY","JUNE",
          "JULY","AUGUST","SEPTEMBER","OCTOBER","NOVEMBER","DECEMBER"]

SUBSYSTEMS = [
    ("PS5-01-01","Fire Water Jockey Pump A/B"),("PS5-01-02","Fire Water Pump A"),
    ("PS5-01-03","Fire Water Pump B"),("PS5-01-04","Fire Water Ring Main"),
    ("PS5-01-05","Foam System"),("PS5-01-06","Deluge System"),
    ("PS5-01-07","Hydrant System"),("PS5-01-08","Drencher System"),
    ("PS5-01-09","Inergen for EITS Building"),("PS5-02-01","Diesel Storage Tank"),
    ("PS5-02-02","Diesel Transfer Pump"),("PS5-02-03","Fuel Gas System"),
    ("PS5-03-01","ESD System"),("PS5-03-02","Fire & Gas Detection"),
    ("PS5-03-03","Flame Detector System"),("PS5-03-23","FZ-E1 to E3 - EITS Building"),
    ("PS5-04-01","DCS System"),("PS5-04-02","SIS System"),
    ("PS5-04-03","33kV ICSS System"),("PS5-04-04","33kV PCS System"),
    ("PS5-05-01","Telecom Network"),("PS5-05-02","PA/GA System"),
    ("PS5-05-03","CCTV System"),("PS5-06-01","AC UPS - EITS Building"),
    ("PS5-06-02","DC UPS - EITS Building"),("PS5-06-03","Lightning Protection"),
    ("PS5-07-01","Chiller System A"),("PS5-07-02","Chiller System B"),
    ("PS5-07-03","AHU System"),("PS5-07-04","FCU System"),
    ("PS5-08-01","Compressor Air System"),("PS5-08-02","Nitrogen Generator"),
    ("PS5-09-01","Transformer 33kV/6.6kV"),("PS5-09-02","Transformer 6.6kV/0.4kV"),
    ("PS5-09-03","Motor Control Center"),("PS5-09-04","0.4kV Switchboard"),
    ("PS5-10-01","Main Pipe Rack"),("PS5-10-02","Utility Pipe Rack"),
    ("PS5-10-03","Process piping A"),("PS5-10-04","Process piping B"),
    ("PS5-10-05","Drain System"),("PS5-10-06","Vent System"),
    ("PS5-11-01","Guard Hut A"),("PS5-11-02","Guard Hut B"),
    ("PS5-11-03","Admin Building"),("PS5-12-01","Cooling Tower"),
    ("PS5-12-02","Water Treatment"),("PS5-13-01","Flare System"),
    ("PS5-60-01","Power Generation Package A"),("PS5-60-04","Power Generation Package B"),
    ("PS5-60-07","Power Generation Package C"),("PS5-60-10","Power Generation Package D"),
    ("PS5-61-01","33kV Switchboard EITS Building"),("PS5-61-02","6.6kV Switchboard LLHT North"),
    ("PS5-61-07","0.4kV Switchboard EITS Building"),("PS5-61-12","EITS 33kV Bld Normal Auxiliary DP"),
    ("PS5-62-01","Cable Tray Main"),("PS5-62-02","Cable Tray Secondary"),
    ("PS5-63-01","Electrical Earthing"),("PS5-63-04","Electrical Earthing EITS Building"),
    ("PS5-82-01","Structural Steel A"),("PS5-82-02","Structural Steel B"),
    ("PS5-82-07","Guard Huts"),("PS5-99-01","Area Completion Building Area"),
    ("PS5-99-02","Area Completion Process Area"),("PS5-99-03","Area Completion Utilities Area"),
]

PUNCH_DESCS = {
    "B":["Missing fire extinguisher signage","Incomplete floor tiling","Damaged ceiling panel","Missing door closer","Cracked wall plaster","Incomplete paint touch-up"],
    "E":["Cable tray not properly supported","Missing earth continuity label","Wrong cable glanding","Incomplete cable termination","Missing electrical hazard sign","Outgoing breaker not labeled"],
    "H":["Duct leakage at joint","Damper not operational","Missing insulation on duct","AHU filter not installed","FCU drain pipe not connected","Thermostat calibration needed"],
    "I":["Transmitter not calibrated","Missing loop diagram","Instrument tubing leak","Pressure gauge reads zero","Control valve actuator fault","Missing instrument tag"],
    "M":["Pump alignment required","Missing rotation arrow","Coupling guard loose","Bearing temperature high","Vibration level exceeded","Motor fan guard missing"],
    "P":["Flange bolt torque incomplete","Pipe support missing","Valve handwheel stiff","Insulation damaged","Missing pipe stress label","Gasket protruding from flange"],
    "S":["Fire extinguisher expired","Emergency shower not tested","Eye wash station blocked","Missing safety sign","Fire alarm panel fault","Deluge valve not primed"],
    "T":["Fiber patch panel unorganized","Missing cable label","CAT6 cable untidied","PA speaker not tested","CCTV camera not focused","Intercom not functioning"],
}

ITR_DESCS = {
    "B":["CQX01 - Steel Structure","CQX02 - Concrete Works","CQX03 - Architectural","CQX04 - Painting & Coating"],
    "E":["CQE01 - HV Cable Installation","CQE02 - LV Cable Installation","CQE03 - Earthing","CQE04 - Transformer","CQE05 - Switchgear"],
    "H":["CQH01 - Ductwork","CQH02 - AHU Installation","CQH03 - FCU Installation","CQH04 - Insulation"],
    "I":["CQI01 - Instrument Tubing","CQI02 - Transmitter Installation","CQI03 - Control Valve","CQI04 - Analyzer Installation"],
    "M":["CQM01 - Pump Installation","CQM02 - Compressor","CQM03 - Heat Exchanger","CQM04 - Vessel","CQM05 - Rotating Equipment"],
    "P":["CQP01 - Pipe Installation","CQP02 - Valve Installation","CQP03 - Flange Management","CQP04 - Hydrotest","CQP05 - Insulation"],
    "S":["CQS01 - Fire Protection","CQS02 - Gas Detection","CQS03 - Deluge System","CQS04 - Foam System"],
    "T":["CQT01 - Cable Installation","CQT02 - Fiber Optic","CQT03 - PA/GA System","CQT04 - CCTV Installation"],
}

TASK_TYPES = ["Conformity Check","Precommissioning","Inspection","Test Pack","Walkdown"]
TASK_STATES = ["To be completed","In progress","Closed","On hold"]
STATUS_OPTIONS = ["Originated","Closed","In Progress","On Hold"]

# ===== STYLE HELPERS =====
HEADER_FILL    = PatternFill(start_color=C_GOLD_DARK,  end_color=C_GOLD_DARK,  fill_type="solid")
HEADER_FONT    = Font(name="Calibri", size=10, bold=True, color=C_WHITE)
SUBHDR_FILL    = PatternFill(start_color=C_GOLD_MED,   end_color=C_GOLD_MED,   fill_type="solid")
SUBHDR_FONT    = Font(name="Calibri", size=9, bold=True, color=C_BROWN_DARK)
TITLE_FONT     = Font(name="Calibri", size=14, bold=True, color=C_GOLD_DARK)
DATA_FONT      = Font(name="Calibri", size=9, color=C_BROWN_DARK)
BOLD_FONT      = Font(name="Calibri", size=10, bold=True, color=C_BROWN_DARK)
GRAND_FONT     = Font(name="Calibri", size=11, bold=True, color=C_WHITE)
NOTE_FONT      = Font(name="Calibri", size=8, italic=True, color=C_BROWN_LIGHT)
THIN_BORDER    = Border(
    left=Side(style='thin', color=C_GOLD_LIGHT),
    right=Side(style='thin', color=C_GOLD_LIGHT),
    top=Side(style='thin', color=C_GOLD_LIGHT),
    bottom=Side(style='thin', color=C_GOLD_LIGHT),
)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ROW_FILL_1 = PatternFill(start_color=C_ROW_1, end_color=C_ROW_1, fill_type="solid")
ROW_FILL_2 = PatternFill(start_color=C_ROW_2, end_color=C_ROW_2, fill_type="solid")
GRAND_FILL  = PatternFill(start_color=C_GOLD_DARK, end_color=C_GOLD_DARK, fill_type="solid")
TOTAL_FILL  = PatternFill(start_color=C_GOLD_LIGHT, end_color=C_GOLD_LIGHT, fill_type="solid")


def rand_date(sy=2026, sm=1, ey=2026, em=12):
    y = random.randint(sy, ey)
    m = random.randint(sm if y == sy else 1, em if y == ey else 12)
    return datetime.date(y, m, random.randint(1, 28))


def gen_sub_desc(sid):
    for s, d in SUBSYSTEMS:
        if s == sid:
            return f"{s} - {d}"
    return f"{sid} - Unknown"


def get_ms_num(p):
    m = {"PS5 - Milestone A":1,"PS5 - Milestone B":1.1,"PS5 - Milestone C":2,
         "PS5 - Milestone D":3,"PS5 - Milestone E":4,"PS5 - Milestone F":5,
         "PS5 - Milestone G":6,"PS5 - Milestone H":7,"PS5 - Milestone I":8}
    return m.get(p, 1)


def get_month(d):
    if d is None: return "#N/A"
    return MONTHS[d.month - 1]


def style_row(ws, r, ncol, idx):
    fill = ROW_FILL_1 if idx % 2 == 0 else ROW_FILL_2
    for c in range(1, ncol + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill = fill
        cell.font = DATA_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def write_grand_row(ws, r, ncol, vals, fonts=None):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.fill = GRAND_FILL
        cell.font = GRAND_FONT if not fonts else fonts[c-1]
        cell.alignment = CENTER
        cell.border = THIN_BORDER


# ===== GENERATE DATA =====
def generate_punch_data(n=2500):
    rows = []
    for i in range(1, n+1):
        cat = random.choice(list(DISCIPLINES.keys()))
        sub_id, sub_name = random.choice(SUBSYSTEMS)
        prio = random.choice(MILESTONES)
        status = random.choices(STATUS_OPTIONS, weights=[10,20,50,20])[0]
        closing = rand_date(2026,8,2026,12) if status=="Closed" else None
        rows.append({
            "id":f"PL-{i:05d}",
            "asset":f"PS5-{random.randint(1,99):02d}-{random.choice(['UR','IK','OB','NER','NI'])}-{random.randint(1000,9999)}{random.choice('ABCD')}",
            "cat":cat,"discipline":cat,
            "desc":random.choice(PUNCH_DESCS.get(cat,["General punch"])),
            "company":random.choice(COMPANIES[:3]),
            "subsystem":gen_sub_desc(sub_id),"closing":closing,"status":status,
            "originated_by":random.choice(["Thao-my, Alexandre","Jean, Pierre","Mohamed, Ali","Smith, John","Kumar, Raj","Wang, Li"]),
            "priority":prio,"priority_num":get_ms_num(prio),
        })
    return rows


def generate_itr_data(n=14000):
    rows = []
    for i in range(1, n+1):
        cat = random.choice(list(DISCIPLINES.keys()))
        sub_id, sub_name = random.choice(SUBSYSTEMS)
        prio = random.choice(MILESTONES)
        state = random.choices(TASK_STATES, weights=[30,20,40,10])[0]
        closing = rand_date(2026,1,2026,12) if state=="Closed" else None
        rows.append({
            "task_id":f"T-{random.randint(3,999):03d}-{i:04d}",
            "asset_tag":f"PS5-{random.randint(1,99):02d}-{random.choice(['OB','UR','IK'])}-{random.randint(1000,9999)}{random.choice('ABCD')}",
            "description":random.choice(ITR_DESCS.get(cat,["General ITR"])),
            "category":"PCOM - Precommissioning" if random.random()>0.3 else "COM - Commissioning",
            "discipline":cat,"discipline_name":DISCIPLINES[cat]["name"],
            "task_type":random.choice(TASK_TYPES),"responsible":"CPP AGI",
            "area":"PS5 - TZA - PS5 - Tanzania","subsystem":gen_sub_desc(sub_id),
            "subsystem_id":sub_id,"closing":closing,"priority":prio,
            "priority_num":get_ms_num(prio),"task_state":state,
            "actual_mh":random.randint(4,48) if state=="Closed" else None,
            "planned_mh":random.randint(4,48),"duration":random.choice([4,8,12,16,24]),
            "month":get_month(closing if closing else rand_date(2026,8,2026,12)),
        })
    return rows


# ====================================================================
# SHEET 1: ITR SUMMARY
# ====================================================================
def build_itr_summary(wb):
    ws = wb.create_sheet("ITR SUMMARY")
    ws.sheet_properties.tabColor = C_GOLD_DARK

    ws.merge_cells('C4:F4')
    ws['C4'] = "Responsible Company : CPP AGI"
    ws['C4'].font = TITLE_FONT
    ws.merge_cells('C5:F5')
    ws['C5'] = "Physical Area : PS5 - Tanzania"
    ws['C5'].font = TITLE_FONT
    ws.merge_cells('C6:G6')
    ws['C6'] = "Pre-commissioning ITRs"
    ws['C6'].font = Font(name="Calibri", size=13, bold=True, color=C_GOLD_DARK)

    headers = ["Discipline","Total","Closed","%","Open","CUT OFF DATE-21-AUGUST-2026"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=7, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    cpp_disc = [("B",71),("H",123),("M",422),("P",2911),("S",172)]
    eit_disc = [("E",4627),("I",4595),("T",1078)]

    row = 8
    ws.cell(row=row, column=1, value="CPP").font = Font(name="Calibri", size=10, bold=True, color=C_GOLD_DARK)
    row += 1

    cpp_t = [0,0,0]
    for code, total in cpp_disc:
        closed = random.randint(int(total*0.05), int(total*0.7))
        pct = round(closed/total, 4)
        ws.cell(row=row, column=1, value=DISCIPLINES[code]["name"]).font = DATA_FONT
        ws.cell(row=row, column=2, value=total).font = DATA_FONT
        ws.cell(row=row, column=3, value=closed).font = DATA_FONT
        ws.cell(row=row, column=4, value=pct).font = DATA_FONT
        ws.cell(row=row, column=5, value=total-closed).font = DATA_FONT
        for c in range(1,6):
            ws.cell(row=row, column=c).alignment = CENTER
            ws.cell(row=row, column=c).border = THIN_BORDER
            ws.cell(row=row, column=c).fill = ROW_FILL_1 if (row-9)%2==0 else ROW_FILL_2
        cpp_t[0]+=total; cpp_t[1]+=closed; cpp_t[2]+=total-closed
        row += 1

    for c, v in enumerate([1, cpp_t[0], cpp_t[1], round(cpp_t[1]/cpp_t[0],4), cpp_t[2]], 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.fill = TOTAL_FILL; cell.font = BOLD_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
    row += 2

    ws.cell(row=row, column=1, value="EIT").font = Font(name="Calibri", size=10, bold=True, color=C_GOLD_DARK)
    row += 1
    eit_t = [0,0,0]
    for code, total in eit_disc:
        closed = random.randint(int(total*0.02), int(total*0.4))
        pct = round(closed/total, 4)
        ws.cell(row=row, column=1, value=DISCIPLINES[code]["name"]).font = DATA_FONT
        ws.cell(row=row, column=2, value=total).font = DATA_FONT
        ws.cell(row=row, column=3, value=closed).font = DATA_FONT
        ws.cell(row=row, column=4, value=pct).font = DATA_FONT
        ws.cell(row=row, column=5, value=total-closed).font = DATA_FONT
        for c in range(1,6):
            ws.cell(row=row, column=c).alignment = CENTER
            ws.cell(row=row, column=c).border = THIN_BORDER
            ws.cell(row=row, column=c).fill = ROW_FILL_1 if row%2==0 else ROW_FILL_2
        eit_t[0]+=total; eit_t[1]+=closed; eit_t[2]+=total-closed
        row += 1

    for c, v in enumerate([1, eit_t[0], eit_t[1], round(eit_t[1]/eit_t[0],4), eit_t[2]], 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.fill = TOTAL_FILL; cell.font = BOLD_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
    row += 2

    g = [cpp_t[0]+eit_t[0], cpp_t[1]+eit_t[1], cpp_t[2]+eit_t[2]]
    write_grand_row(ws, row, 5, ["Grand Total", g[0], g[1], round(g[1]/g[0],4), g[2]])
    row += 2
    ws.cell(row=row, column=1, value=f"*Note: Generated from Hexagon data as of- {TODAY_STR} at 7.00 AM").font = NOTE_FONT

    for c in range(1, 10):
        ws.column_dimensions[get_column_letter(c)].width = 18
    return {"cpp":cpp_t, "eit":eit_t, "grand":g}


# ====================================================================
# SHEET 2: PUNCH SUMMARY
# ====================================================================
def build_punch_summary(wb):
    ws = wb.create_sheet("PUNCH SUMMERY")
    ws.sheet_properties.tabColor = C_ACCENT_RED

    ws.merge_cells('C4:F4')
    ws['C4'] = "PS5  Punch Summary"
    ws['C4'].font = TITLE_FONT

    cat_names = ["A Punch Summary","B Punch Summary","C Punch Summary","0 Punch Summary"]
    cat_cols = [3,9,15,21]

    for name, col in zip(cat_names, cat_cols):
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col+2)
        for cc in range(col, col+3):
            ws.cell(row=5, column=cc).fill = HEADER_FILL
            ws.cell(row=5, column=cc).border = THIN_BORDER
        ws.cell(row=5, column=col, value=name).font = HEADER_FONT
        ws.cell(row=5, column=col).alignment = CENTER

    for col_start in cat_cols:
        for i, h in enumerate(["Discipline","Total","Closed","Open"]):
            cell = ws.cell(row=6, column=col_start+i, value=h)
            cell.fill = SUBHDR_FILL; cell.font = SUBHDR_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    disc_list = ["Building","Electrical","HVAC","Instrumentation","Mechanical","Piping & Vessel","Safety","Telecom","Miscellaneous"]
    for d_idx, disc in enumerate(disc_list):
        r = 7 + d_idx
        for ci, col in enumerate(cat_cols):
            total = random.randint(0,600) if disc!="Miscellaneous" else 0
            closed = random.randint(0,total)
            for j, v in enumerate([disc, total, closed, total-closed]):
                cell = ws.cell(row=r, column=col+j, value=v)
                cell.font = DATA_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
                cell.fill = ROW_FILL_1 if d_idx%2==0 else ROW_FILL_2

    r = 7 + len(disc_list)
    for ci, col in enumerate(cat_cols):
        gt = sum((ws.cell(row=7+d, column=col+1).value or 0) for d in range(len(disc_list)))
        gc = sum((ws.cell(row=7+d, column=col+2).value or 0) for d in range(len(disc_list)))
        vals = ["Grand Total", gt, gc, gt-gc]
        for j, v in enumerate(vals):
            cell = ws.cell(row=r, column=col+j, value=v)
            cell.fill = GRAND_FILL; cell.font = GRAND_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    r += 2
    ws.cell(row=r, column=1, value=f"*Note: Generated from Hexagon data as of - {TODAY_STR} at 7.00 AM").font = NOTE_FONT
    for c in range(1, 26):
        ws.column_dimensions[get_column_letter(c)].width = 14


# ====================================================================
# SHEET 3: DETAILED PUNCH LIST
# ====================================================================
def build_detailed_punch(wb, data):
    ws = wb.create_sheet("DETAILED PUNCH LIST")
    ws.sheet_properties.tabColor = C_GOLD_MED
    headers = ["Punchlist ID","Asset (Name/Tag)","Asset Pack","CAT","Discipline (Name)","Description",
               "Scheduling - Responsible Company (Name)","Systemization - Subsystem (Summary)",
               "Workflow - Closing Date","Status","Originated by","PRIORITY","PRIORITY_NUM"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
    for i, p in enumerate(data):
        r = i + 2
        vals = [p["id"],p["asset"],None,p["cat"],p["discipline"],p["desc"],p["company"],
                p["subsystem"],p["closing"],p["status"],p["originated_by"],p["priority"],p["priority_num"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = DATA_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
            cell.fill = ROW_FILL_1 if i%2==0 else ROW_FILL_2
    for c in range(1,14):
        ws.column_dimensions[get_column_letter(c)].width = 20
    ws.column_dimensions['F'].width = 45


# ====================================================================
# SHEET 4: DETAILED ITR LIST
# ====================================================================
def build_detailed_itr(wb, data):
    ws = wb.create_sheet("DETAILED ITR LIST")
    ws.sheet_properties.tabColor = C_ACCENT_GREEN
    headers = ["Task ID","Asset - Tag","Description","Category (Summary)","Discipline",
               "Task Type (Name)","Discipline (Summary)","Responsible Company (Summary)",
               "Physical Location - Area (Summary)","Systemization - Subsystem (Summary)",
               "Closing Date","Subsystem ID","Priority","Priority_NUM","Task State",
               "Actual MH","Planned MH","MONTH"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
    for i, t in enumerate(data):
        r = i + 2
        vals = [t["task_id"],t["asset_tag"],t["description"],t["category"],t["discipline"],
                t["task_type"],t["discipline_name"],t["responsible"],t["area"],t["subsystem"],
                t["closing"],t["subsystem_id"],t["priority"],t["priority_num"],t["task_state"],
                t["actual_mh"],t["planned_mh"],t["month"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = DATA_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
            cell.fill = ROW_FILL_1 if i%2==0 else ROW_FILL_2
    for c in range(1,19):
        ws.column_dimensions[get_column_letter(c)].width = 22


# ====================================================================
# SHEET 5: P1&P2 (NEW)
# ====================================================================
def build_p1p2(wb):
    ws = wb.create_sheet("P1&P2")
    ws.sheet_properties.tabColor = C_ACCENT_RED

    ws.merge_cells('A1:Q1')
    ws['A1'] = "PRESURE TEST AND REINSTATEMENT-PRIORITY 1 & 2"
    ws['A1'].font = TITLE_FONT
    ws['A1'].alignment = CENTER

    h2 = ["Pipe Pack ID","SUBSYSTEM WITH ID","PRIORITY"]
    h2_cols = [(1,3),(4,5),(6,7)]
    sub_h = ["PRESSURE TEST STATUS","","REINSTATEMENT STATUS","","GALT","","PAINTING LINE","","PRESONAL Protection","","STATUS"]
    sub_cols = [8,9,10,11,12,13,14,15,16,17]

    for c, h in enumerate(["Pipe Pack ID","SUBSYSTEM WITH ID","PRIORITY",
                           "PRESSURE TEST STATUS","","REINSTATEMENT STATUS","","GALT","","PAINTING LINE","","PRESONAL PROTECTION","","STATUS"], 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    for c, h in enumerate(["","","","ITR NUMBER","STATUS","ITR NUMBER","STATUS","TAG","ITR","STATUS","ITR NUMBER","TAG","STATUS","ITR NUMBER","TAG","STATUS","STATUS"], 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.fill = SUBHDR_FILL; cell.font = SUBHDR_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    for i in range(90):
        r = i + 4
        sub_id, sub_name = random.choice(SUBSYSTEMS)
        pipe_pack = f"EACOP-{sub_id}-{random.randint(100,999)}"
        prio = random.choice(["Priority-1","Priority-2"])
        pt_status = random.choice(["Closed","Open","N/A"])
        rt_status = random.choice(["Closed","Started (not Completed)","N/A"])
        galt_status = random.choice(["Closed","N/A"])
        overall = random.choice(["INT.WALKDOWN COMPLETED","PENDING","IN PROGRESS","COMPLETED"])

        vals = [pipe_pack, gen_sub_desc(sub_id), prio,
                f"T-{random.randint(200,300)}-{random.randint(100,999)}", pt_status,
                f"T-{random.randint(200,300)}-{random.randint(100,999)}", rt_status,
                f"{sub_id}_GALT", f"T-{random.randint(200,300)}-{random.randint(100,999)}", galt_status,
                f"T-{random.randint(200,300)}-{random.randint(100,999)}", f"{sub_id}_SS_ALL_LINES", rt_status,
                "#N/A","#N/A","#N/A", overall]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = DATA_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
            cell.fill = ROW_FILL_1 if i%2==0 else ROW_FILL_2

    for c in range(1,18):
        ws.column_dimensions[get_column_letter(c)].width = 18


# ====================================================================
# SHEET 6: MILESTONE_12 AUG 2026 (NEW)
# ====================================================================
def build_milestone_12aug(wb):
    ws = wb.create_sheet("MILESTONE_12 AUG 2026")
    ws.sheet_properties.tabColor = C_BROWN_MED

    headers = ["SUBSYSTEM","SUBSYSTEM DESCRIPTION","SUBSYSTEM DESCRIPTIONS","Priority",
               "MONTH","mantrac","PRIORITY UPDATED-08-08-2026","PG A TO G","PG/COMMON",
               "RFSU","M & P DATES","EIT DATES","RFC DATES","PRIORITY UPDATED-11-07-2026"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    for i, (sid, sname) in enumerate(SUBSYSTEMS):
        r = i + 2
        prio = random.choice(MILESTONES)
        month = random.choice(MONTHS + ["#N/A"])
        mantrac = random.choice(["yes","no"])
        rfsu = random.choice(["Yes","No","#N/A","complete"])
        m_p = random.choice(["complete","#N/A",""])
        eit = random.choice(["complete","#N/A","N/A",""])
        rfc = random.choice(["complete","#N/A","N/A",""])
        pg = random.choice([None, "#N/A"])
        pg_common = random.choice(["Yes","#N/A"])
        prio_updated = random.choice([1,2,3,4,5,6,7,8,9,10,"EITS",None])

        vals = [sid, sname, gen_sub_desc(sid), prio, month, mantrac, prio_updated,
                pg, pg_common, rfsu, m_p, eit, rfc]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = DATA_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
            cell.fill = ROW_FILL_1 if i%2==0 else ROW_FILL_2

    for c in range(1,15):
        ws.column_dimensions[get_column_letter(c)].width = 22


# ====================================================================
# SHEET 7: MILESTONE & DIS.
# ====================================================================
def build_milestone(wb):
    ws = wb.create_sheet("MILESTONE & DIS.")
    ws.sheet_properties.tabColor = C_BROWN_LIGHT

    ws.merge_cells('B3:E3')
    ws['B3'] = "Responsible Company : CPP AGI"
    ws['B3'].font = TITLE_FONT
    ws.merge_cells('B4:E4')
    ws['B4'] = "Physical Area : PS5 - Tanzania"
    ws['B4'].font = TITLE_FONT
    ws.merge_cells('B5:E5')
    ws['B5'] = "Pre-commissioning ITRs"
    ws['B5'].font = Font(name="Calibri", size=13, bold=True, color=C_GOLD_DARK)

    for c, h in enumerate(["Discipline","Total","Closed","%","Open"], 1):
        cell = ws.cell(row=7, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    row = 8
    for code, info in DISCIPLINES.items():
        total = random.randint(50, 3000)
        closed = random.randint(0, total)
        for c, v in enumerate([info["name"], total, closed, round(closed/total,4), total-closed], 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = DATA_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
            cell.fill = ROW_FILL_1 if (row-8)%2==0 else ROW_FILL_2
        row += 1
    for c in range(1,8):
        ws.column_dimensions[get_column_letter(c)].width = 20


# ====================================================================
# SHEET 8: RFC PROGRESS
# ====================================================================
def build_rfc_progress(wb):
    ws = wb.create_sheet("RFC PROGRESS")
    ws.sheet_properties.tabColor = C_GOLD_MED

    headers = ["SUBSYSTEM ID WITH DESCRIPTION","PRIORITY","RFC DATE","MILESTONE","TOTAL %",
               "B-BUILDING T","B-CLOSED","B-BAL","H-HVAC T","H-CLOSED","H-BAL",
               "M-MECHANICAL T","M-CLOSED","M-BAL","P-PIPING T","P-CLOSED","P-BAL",
               "S-SAFETY T","S-CLOSED","S-BAL","TOTAL ITRs","TOTAL CLOSED","TOTAL BAL","ITR %",
               "BLOCKING POINTS REMARKS"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    for i, (sid, sn) in enumerate(SUBSYSTEMS):
        r = i + 2
        prio = random.choice(MILESTONES)
        rfc_date = rand_date(2026,3,2026,12)
        totals = {}
        for disc in ["B","H","M","P","S"]:
            t = random.randint(0,50)
            c = random.randint(0,t)
            totals[disc] = (t,c,t-c)
        gt = sum(v[0] for v in totals.values())
        gc = sum(v[1] for v in totals.values())
        pct = round(gc/gt,4) if gt>0 else 0
        remark = random.choice(["","WAITING FOR MATERIALS","ALIGNMENT TOOL REQUIRED","EACOP SUPPORT PENDING","MANTRAC SCOPE"])

        ws.cell(row=r, column=1, value=gen_sub_desc(sid)).font = DATA_FONT
        ws.cell(row=r, column=2, value=get_ms_num(prio)).font = DATA_FONT
        ws.cell(row=r, column=3, value=rfc_date).font = DATA_FONT
        ws.cell(row=r, column=4, value=prio).font = DATA_FONT
        ws.cell(row=r, column=5, value=pct).font = DATA_FONT
        col = 6
        for disc in ["B","H","M","P","S"]:
            ws.cell(row=r, column=col, value=totals[disc][0]).font = DATA_FONT
            ws.cell(row=r, column=col+1, value=totals[disc][1]).font = DATA_FONT
            ws.cell(row=r, column=col+2, value=totals[disc][2]).font = DATA_FONT
            col += 3
        ws.cell(row=r, column=21, value=gt).font = DATA_FONT
        ws.cell(row=r, column=22, value=gc).font = DATA_FONT
        ws.cell(row=r, column=23, value=gt-gc).font = DATA_FONT
        ws.cell(row=r, column=24, value=pct).font = DATA_FONT
        ws.cell(row=r, column=25, value=remark).font = DATA_FONT

        for c in range(1, 26):
            ws.cell(row=r, column=c).alignment = CENTER
            ws.cell(row=r, column=c).border = THIN_BORDER
            ws.cell(row=r, column=c).fill = ROW_FILL_1 if i%2==0 else ROW_FILL_2
    for c in range(1,26):
        ws.column_dimensions[get_column_letter(c)].width = 16


# ====================================================================
# SHEET 9: M-P PROGRESS (NEW)
# ====================================================================
def build_mp_progress(wb):
    ws = wb.create_sheet("M-P PROGRESS")
    ws.sheet_properties.tabColor = C_BROWN_MED

    headers = ["SUBSYSTEM ID WITH DESCRIPTION","PRIORITY","NEW RFC DATE(B,H,M,P,S)","MILESTONE",
               "B-TOTAL","B-CLOSED","B-BAL","B-%",
               "H-TOTAL","H-CLOSED","H-BAL","H-%",
               "M-TOTAL","M-CLOSED","M-BAL","M-%",
               "P-TOTAL","P-CLOSED","P-BAL","P-%",
               "S-TOTAL","S-CLOSED","S-BAL","S-%",
               "TOTAL ITRs","TOTAL CLOSED","BALANCE","%",
               "PRIORITY 1&2","BLOCKING ITRs","BLOCKING POINTS-CPY","BLOCKING POINTS-CPP","REMARKS"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    for i, (sid, sn) in enumerate(SUBSYSTEMS):
        r = i + 2
        prio = random.choice(MILESTONES)
        rfc_date = rand_date(2026,3,2026,12)
        disc_data = {}
        for disc in ["B","H","M","P","S"]:
            t = random.randint(0,60)
            c = random.randint(0,t)
            disc_data[disc] = (t,c,t-c,round(c/t,4) if t>0 else 0)

        gt = sum(v[0] for v in disc_data.values())
        gc = sum(v[1] for v in disc_data.values())
        pct = round(gc/gt,4) if gt>0 else 0
        blocking = random.randint(0,15)
        remark = random.choice(["","WAITING FOR MATERIALS","ALIGNMENT TOOL REQUIRED","EACOP SCOPE","MANTRAC SCOPE"])

        vals = [gen_sub_desc(sid), get_ms_num(prio), rfc_date, prio]
        for disc in ["B","H","M","P","S"]:
            vals.extend(list(disc_data[disc]))
        vals.extend([gt, gc, gt-gc, pct, "", blocking, "", "", remark])

        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = DATA_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
            cell.fill = ROW_FILL_1 if i%2==0 else ROW_FILL_2

    for c in range(1,33):
        ws.column_dimensions[get_column_letter(c)].width = 14


# ====================================================================
# SHEET 10: EIT PROGRESS
# ====================================================================
def build_eit_progress(wb):
    ws = wb.create_sheet("EIT PROGRESS")
    ws.sheet_properties.tabColor = C_ACCENT_BLUE

    headers = ["SUBSYSTEM DESCRIPTION","PRIORITY","NEW RFC DATE EIT","MILESTONE",
               "E-ELECTRICAL T","E-CLOSED","E-REMAIN","E-%",
               "I-INSTRUMENTATION T","I-CLOSED","I-REMAIN","I-%",
               "T-TELECOM T","T-CLOSED","T-REMAIN","T-%",
               "TOTAL","CLOSED","BALANCE","%",
               "BLOCKING POINTS-CPY","BLOCKING POINTS-CPP","REMARKS"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    for i, (sid, sn) in enumerate(SUBSYSTEMS):
        r = i + 2
        prio = random.choice(MILESTONES)
        rfc_date = rand_date(2026,3,2026,12)

        e_t=random.randint(0,100); e_c=random.randint(0,e_t); e_r=e_t-e_c; e_pct=round(e_c/e_t,4) if e_t>0 else 0
        i_t=random.randint(0,80);  i_c=random.randint(0,i_t);  i_r=i_t-i_c;  i_pct=round(i_c/i_t,4) if i_t>0 else 0
        t_t=random.randint(0,30);  t_c=random.randint(0,t_t);  t_r=t_t-t_c;  t_pct=round(t_c/t_t,4) if t_t>0 else 0
        total=e_t+i_t+t_t; closed=e_c+i_c+t_c; balance=total-closed; pct=round(closed/total,4) if total>0 else 0
        remark = random.choice(["","EACOP MATERIAL PENDING","TRANSFORMER NOT ARRIVED","CABLE DRUM SHORTAGE","EARTHING PENDING"])

        vals = [gen_sub_desc(sid), get_ms_num(prio), rfc_date, prio,
                e_t,e_c,e_r,e_pct, i_t,i_c,i_r,i_pct, t_t,t_c,t_r,t_pct,
                total,closed,balance,pct, "", "", remark]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = DATA_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
            cell.fill = ROW_FILL_1 if i%2==0 else ROW_FILL_2
    for c in range(1,24):
        ws.column_dimensions[get_column_letter(c)].width = 18


# ====================================================================
# SHEET 11: EIT - PROGRESS (NEW - placeholder like original)
# ====================================================================
def build_eit_progress_empty(wb):
    ws = wb.create_sheet("EIT - PROGRESS")
    ws.sheet_properties.tabColor = C_GOLD_LIGHT
    ws.cell(row=1, column=1, value="EIT - PROGRESS").font = TITLE_FONT


# ====================================================================
# SHEET 12: RFC DATES OLD (NEW)
# ====================================================================
def build_rfc_dates_old(wb):
    ws = wb.create_sheet("RFC DATES OLD")
    ws.sheet_properties.tabColor = C_BROWN_LIGHT

    headers = ["S.no","Process","SubSystem","SubSystem Description","SubSystem ID With Description",
               "Milestone","Required for Start-up","CPP AGI SCOPE","KENT Scope","EITS Scope",
               "Structure","Mechanical","Piping","Paint/Coat/Insul","EIT",
               "Final CPP RFC Date","a","b","PRIORITY-08-08-2026"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER

    processes = [
        "PS5-01 - Fire Fighting Generation / Distribution",
        "PS5-02 - Diesel Storage & Transfer",
        "PS5-03 - ESD and Fire & gas",
        "PS5-04 - Surface Control & Supervision",
        "PS5-05 - Telecommunication",
        "PS5-06 - AC/DC UPS",
        "PS5-07 - HVAC",
        "PS5-08 - Instrument Air",
        "PS5-09 - Power Distribution",
        "PS5-10 - Piping",
        "PS5-11 - Buildings",
        "PS5-12 - Utilities",
        "PS5-13 - Flare System",
        "PS5-60 - Power Generation",
        "PS5-61 - HV Switchgear",
        "PS5-62 - Cable Tray",
        "PS5-63 - Earthing",
        "PS5-82 - Structural Steel",
    ]
    scopes = ["CPP AGI","KENT","EITS",None]

    for i in range(280):
        r = i + 2
        sub_id, sub_name = random.choice(SUBSYSTEMS)
        proc = random.choice(processes)
        prio = random.choice(MILESTONES)
        startup = random.choice(["Yes","No"])
        cpp_scope = random.choice(scopes)
        kent_scope = random.choice(scopes)
        eits_scope = random.choice(scopes)

        struct_date = random.choice([rand_date(2026,3,2026,12), None, "-"])
        mech_date = random.choice([rand_date(2026,3,2026,12), None, "-"])
        pipe_date = random.choice([rand_date(2026,3,2026,12), None, "-"])
        paint_date = random.choice([rand_date(2026,3,2026,12), None, "-"])
        eit_date = random.choice([rand_date(2026,3,2026,12), "EITS TBC", None, "-"])
        final_date = random.choice([rand_date(2026,5,2026,12), "EITS TBC", None])
        priority = random.choice([None, "Priority-1", "Priority-2"])

        vals = [i+1, proc, sub_id, sub_name, gen_sub_desc(sub_id), prio, startup,
                cpp_scope, kent_scope, eits_scope,
                struct_date, mech_date, pipe_date, paint_date, eit_date,
                final_date, None, False, priority]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = DATA_FONT; cell.alignment = CENTER; cell.border = THIN_BORDER
            cell.fill = ROW_FILL_1 if i%2==0 else ROW_FILL_2

    for c in range(1,20):
        ws.column_dimensions[get_column_letter(c)].width = 20


# ====================================================================
# MAIN
# ====================================================================
def main():
    print("Generating dummy data...")
    punch_data = generate_punch_data(2500)
    itr_data = generate_itr_data(14000)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    print("1/12  ITR SUMMARY...")
    summary = build_itr_summary(wb)
    print("2/12  PUNCH SUMMERY...")
    build_punch_summary(wb)
    print("3/12  DETAILED PUNCH LIST...")
    build_detailed_punch(wb, punch_data)
    print("4/12  DETAILED ITR LIST...")
    build_detailed_itr(wb, itr_data)
    print("5/12  P1&P2...")
    build_p1p2(wb)
    print("6/12  MILESTONE_12 AUG 2026...")
    build_milestone_12aug(wb)
    print("7/12  MILESTONE & DIS...")
    build_milestone(wb)
    print("8/12  RFC PROGRESS...")
    build_rfc_progress(wb)
    print("9/12  M-P PROGRESS...")
    build_mp_progress(wb)
    print("10/12 EIT PROGRESS...")
    build_eit_progress(wb)
    print("11/12 EIT - PROGRESS...")
    build_eit_progress_empty(wb)
    print("12/12 RFC DATES OLD...")
    build_rfc_dates_old(wb)

    excel_name = f"PS-5 COMPLETIONS DPR SUMMERY - {DATE_FILE}.xlsx"
    excel_path = os.path.join(OUTPUT_DIR, excel_name)
    wb.save(excel_path)
    print(f"\nExcel saved: {excel_path}")

    # JSON for dashboard
    dashboard_data = {
        "date": TODAY_STR,
        "date_file": DATE_FILE,
        "excel_name": excel_name,
        "itr_summary": summary,
        "punch_by_cat": {},
        "disciplines": {},
        "subsystems_count": len(SUBSYSTEMS),
    }
    for p in punch_data:
        cat = p["cat"]
        if cat not in dashboard_data["punch_by_cat"]:
            dashboard_data["punch_by_cat"][cat] = {"total":0,"closed":0,"open":0}
        dashboard_data["punch_by_cat"][cat]["total"] += 1
        if p["status"]=="Closed":
            dashboard_data["punch_by_cat"][cat]["closed"] += 1
        else:
            dashboard_data["punch_by_cat"][cat]["open"] += 1

    for t in itr_data:
        disc = t["discipline"]
        if disc not in dashboard_data["disciplines"]:
            dashboard_data["disciplines"][disc] = {"total":0,"closed":0,"in_progress":0,"pending":0}
        dashboard_data["disciplines"][disc]["total"] += 1
        if t["task_state"]=="Closed":
            dashboard_data["disciplines"][disc]["closed"] += 1
        elif t["task_state"]=="In progress":
            dashboard_data["disciplines"][disc]["in_progress"] += 1
        else:
            dashboard_data["disciplines"][disc]["pending"] += 1

    json_path = os.path.join(OUTPUT_DIR, "dashboard_data.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(dashboard_data, f, indent=2, default=str)
    print(f"JSON saved: {json_path}")
    print("\nDone! 12/12 sheets generated.")


if __name__ == "__main__":
    main()
