import openpyxl
wb = openpyxl.load_workbook(r'C:\Users\mylap\OneDrive\Desktop\dashboard\PS-5 COMPLETIONS DPR SUMMERY - 19-08-26.xlsx', data_only=True)

# Unique subsystem IDs from ITR list column 11 (Subsystem ID) - full scan
ws = wb['DETAILED ITR LIST']
sub_ids = set()
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=12, max_col=12, values_only=True):
    if row[0]:
        sub_ids.add(str(row[0]).strip())
print(f'Unique Subsystem IDs in ITR list: {len(sub_ids)}')
for s in sorted(sub_ids)[:20]:
    print(f'  {s}')
print('...')
for s in sorted(sub_ids)[-5:]:
    print(f'  {s}')

# Also check unique subsystem summaries (col 9)
sub_summaries = set()
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=10, max_col=10, values_only=True):
    if row[0]:
        sub_summaries.add(str(row[0]).strip())
print(f'\nUnique Subsystem Summaries in ITR list: {len(sub_summaries)}')
for s in sorted(sub_summaries)[:10]:
    print(f'  {s}')
