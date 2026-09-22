import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE = r"C:\Users\mylap\OneDrive\Desktop\dashboard"
DATA_FILE = os.path.join(BASE, "data.json")
OUTPUT = os.path.join(os.environ['USERPROFILE'], 'OneDrive', 'Desktop', 'EIT Cables.xlsx')

with open(DATA_FILE, 'r', encoding='utf-8') as f:
    data = json.load(f)

wb = Workbook()
ws = wb.active
ws.title = "EIT Cables"

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

headers = ["#", "Subsystem", "Asset Tag", "From", "To", "PDF Link", "RFI", "RFI Status", "RFI PDF", "All RFI PDFs"]
col_widths = [5, 40, 25, 45, 45, 50, 20, 15, 50, 60]

for c, (h, w) in enumerate(zip(headers, col_widths), 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border
    ws.column_dimensions[cell.column_letter].width = w

ws.auto_filter.ref = f"A1:J{len(data)+1}"
ws.freeze_panes = "A2"

for i, d in enumerate(data, 1):
    row = i + 1
    vals = [
        i,
        d.get("subsystem", ""),
        d.get("asset_tag", ""),
        d.get("from", ""),
        d.get("to", ""),
        d.get("link", ""),
        d.get("rfi", ""),
        d.get("rfi_status", ""),
        d.get("rfi_pdf", ""),
        ", ".join(d.get("all_rfi_pdfs", []))
    ]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.border = thin_border
        if c == 6 and v:
            cell.hyperlink = v
            cell.font = Font(color="1565C0", underline="single")
        if c == 9 and v:
            cell.hyperlink = v
            cell.font = Font(color="1565C0", underline="single")

wb.save(OUTPUT)
print(f"DONE: {len(data)} cables exported to {OUTPUT}")
