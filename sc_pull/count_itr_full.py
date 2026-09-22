import json
import os
import sys
import time
from collections import Counter

import openpyxl
from playwright.sync_api import sync_playwright

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
WS_DIR = os.path.dirname(BASE_DIR)
DATA_DIR = os.path.join(BASE_DIR, "data")
PROFILE_DIR = os.path.join(BASE_DIR, "profile")

SWITCHBOARD_URL = (
    "https://wly04-sc.intergraphsmartcloud.com/ISC/Tools/vDashboardsUsers/Switchboard.htm"
)

CPP_CID = 1021
PAGE = 2000
MAX_ROWS = 200000

TASK_FIELDS = [
    "ID", "TaskName", "AssetTag", "AssetDescription", "Description", "TaskCategorySummary",
    "TaskDiscipline", "TaskDisciplineSummary", "Discipline", "TaskType", "ResponsibleCompany",
    "ResponsibleCompanyID", "ProcessBreakdown", "ProcessBreakdownIdentifier",
    "ProcessBreakdownUp3Identifier", "PhysicalLocationUp1Summary", "ApprovedDate", "ApprovedBy",
    "TaskState", "AssetType", "LoopName", "TaskModelName", "TaskPrioritySummary", "DocumentDescription",
]

GET_JS = """((P) => {
    return new Promise((res) => {
        let done = false;
        const t = setTimeout(() => { if (!done) { done = true; res({timeout: true}); } }, P.to);
        dalc.Get(P.v, P.f, P.r || [], {Distinct:false, FirstRecordOrdinal: P.off, NumberOfRecords:P.n}, "Flat",
            (ok) => {
                if (done) return; done = true; clearTimeout(t);
                let cols = [], rows = [], total = null;
                const vals = (ok && ok.Values) || [];
                if (vals.length) cols = Object.keys(vals[0]).filter(k=>!k.startsWith('__'));
                rows = JSON.parse(JSON.stringify(vals || []));
                if (ok && ok.Request) {
                    const tr = ok.Request.TotalRecords;
                    if (typeof tr !== 'undefined' && tr !== null) total = tr;
                }
                res({cols: cols, rows: rows, total: total, n: rows.length});
            },
            (err) => { if (!done) { done = true; clearTimeout(t); res({err: String((err && err.message) || err)}); } }, false);
    });
})"""


def call(page, view, fields, restrictions=None, n=PAGE, off=0, to=150000):
    return page.evaluate("(%s)(%s)" % (GET_JS, json.dumps({
        "v": view, "f": fields, "r": restrictions or [], "n": n, "off": off, "to": to,
    })))


def count_products(page):
    try:
        return page.evaluate(
            """() => (typeof switchboardConfiguration!=="undefined" && switchboardConfiguration.Products
            && switchboardConfiguration.Products.length) ? switchboardConfiguration.Products.length : -1"""
        )
    except Exception:
        return -1


def wait_products(page, timeout_sec):
    start = time.time()
    while time.time() - start < timeout_sec:
        if count_products(page) >= 0:
            return True
        try:
            page.wait_for_timeout(2000)
        except Exception:
            pass
    return False


def disc_code(v):
    if v is None:
        return None
    s = str(v).strip()
    if s in ("E", "I", "T", "M", "S", "P", "B", "H", "PL", "IT", "O", "Z", "V", "N"):
        return s
    for pre in ("E -", "I -", "T -", "M -", "S -", "P -", "B -", "H -", "PL -", "IT -"):
        if s.startswith(pre):
            return pre[0]
    return s[:2].upper() if len(s) >= 2 else s.upper()


def excel_reference():
    p = os.path.join(WS_DIR, "ovTasks_TestsPlanned_1369.xlsx")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb["Exported from SC"]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    idx = {h: i for i, h in enumerate(header)}
    cpp_tasks = 0
    cpp_eit = 0
    cpp_eit_closed = 0
    cpp_eit_ps5 = 0
    cpp_eit_ps5_closed = 0
    by_disc = Counter()
    for r in it:
        if str(r[idx["Responsible Company (Summary)"]]) != "CPP AGI":
            continue
        cpp_tasks += 1
        disc = disc_code(r[idx["Discipline (Summary)"]])
        if disc not in ("E", "I", "T"):
            continue
        cpp_eit += 1
        by_disc[disc] += 1
        st = str(r[idx["Task State"]] or "")
        tag = str(r[idx["Asset - Tag"]] or "")
        if st == "Closed":
            cpp_eit_closed += 1
        if tag.startswith("PS5") or "PS5" in str(r[idx["Systemization - Subsystem (Summary)"]]) or "PS5" in str(r[idx["Physical Location - Area (Summary)"]]):
            cpp_eit_ps5 += 1
            if st == "Closed":
                cpp_eit_ps5_closed += 1
    return {
        "cpp_tasks": cpp_tasks, "cpp_eit": cpp_eit, "cpp_eit_closed": cpp_eit_closed,
        "cpp_eit_ps5": cpp_eit_ps5, "cpp_eit_ps5_closed": cpp_eit_ps5_closed,
        "by_disc": dict(by_disc),
    }


def main():
    os.makedirs(DATA_DIR, exist_ok=True)

    ref = None
    try:
        ref = excel_reference()
        print("EXCEL REFERENCE:", json.dumps(ref, ensure_ascii=False), flush=True)
    except Exception as e:
        print("excel ref err:", str(e)[:200], flush=True)

    with sync_playwright() as pw:
        ctx = pw.chromium.launch_persistent_context(
            PROFILE_DIR, channel="chrome", headless=False,
            args=["--start-maximized"], viewport=None,
        )
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        try:
            page.goto(SWITCHBOARD_URL, wait_until="domcontentloaded", timeout=60000)
        except Exception:
            pass
        if not wait_products(page, timeout_sec=120):
            print("NO SESSION (long poll)", flush=True)
            for i in range(60):
                if count_products(page) >= 0:
                    break
                page.wait_for_timeout(5000)
        if count_products(page) < 0:
            page.bring_to_front()
            if not wait_products(page, timeout_sec=900):
                ctx.close()
                return 1
        print("LOGGED IN", flush=True)

        restrictions = [{"FieldName": "ResponsibleCompanyID", "Term": "=", "ValueCollection": [CPP_CID]}]
        chunks_dir = os.path.join(DATA_DIR, "cpp_chunks")
        os.makedirs(chunks_dir, exist_ok=True)
        all_rows = []
        start_off = 0
        while os.path.exists(os.path.join(chunks_dir, "chunk_%d.json" % start_off)):
            with open(os.path.join(chunks_dir, "chunk_%d.json" % start_off), "r", encoding="utf-8") as f:
                all_rows.extend(json.load(f))
            start_off += PAGE
        print("resume from off=%d accumulated=%d" % (start_off, len(all_rows)), flush=True)
        off = start_off
        done = False
        while not done:
            try:
                r = call(page, "vTasks_TestsPlanned", TASK_FIELDS, restrictions, off=off)
            except Exception as e:
                print("exception off=%d: %s" % (off, str(e)[:200]), flush=True)
                break
            n = r.get("n", 0)
            if r.get("timeout"):
                print("timeout at off=%d" % off, flush=True)
                break
            if r.get("err"):
                print("err at off=%d: %s" % (off, r["err"][:150]), flush=True)
                break
            page_rows = r.get("rows", [])
            all_rows.extend(page_rows)
            with open(os.path.join(chunks_dir, "chunk_%d.json" % off), "w", encoding="utf-8") as f:
                json.dump(page_rows, f, ensure_ascii=False)
            print("pulled off=%d n=%d accumulated=%d" % (off, n, len(all_rows)), flush=True)
            if n < PAGE:
                done = True
            off += PAGE
            if len(all_rows) >= MAX_ROWS:
                print("MAXROWS", flush=True)
                done = True

        print("TOTAL CPP AGI tasks pulled:", len(all_rows), flush=True)

        by_disc = Counter()
        closed = Counter()
        open_count = Counter()
        station = Counter()
        station_closed = Counter()
        ps5_eit = 0
        ps5_eit_closed = 0
        for r in all_rows:
            disc = disc_code(r.get("TaskDisciplineSummary") or r.get("TaskDiscipline"))
            if disc not in ("E", "I", "T"):
                continue
            by_disc[disc] += 1
            st = str(r.get("TaskState") or "")
            loc = str(r.get("PhysicalLocationUp1Summary") or "")
            pb = str(r.get("ProcessBreakdown") or "")
            tag = str(r.get("AssetTag") or "")
            is_ps5 = tag.startswith("PS5") or pos8(loc) or pos8(pb)
            if is_ps5:
                ps5_eit += 1
            stname = st if st else "None"
            if stname == "Closed":
                closed[disc] += 1
                if is_ps5:
                    ps5_eit_closed += 1
            else:
                open_count[disc] += 1
            stkey = loc.split(" - ")[0] if loc else pb.split(" - ")[0]
            station[stkey] += 1
            if stname == "Closed":
                station_closed[stkey] += 1

        counts = {
            "total_pulled": len(all_rows),
            "eit_total": sum(by_disc.values()),
            "eit_by_discipline": dict(by_disc),
            "eit_closed": sum(closed.values()),
            "eit_closed_by_discipline": dict(closed),
            "eit_open": sum(open_count.values()),
            "eit_open_by_discipline": dict(open_count),
            "eit_ps5": ps5_eit,
            "eit_ps5_closed": ps5_eit_closed,
            "by_station_total": dict(station),
            "by_station_closed": dict(station_closed),
        }
        print("\n=== LIVE COUNTS ===", flush=True)
        print(json.dumps(counts, ensure_ascii=False, indent=2), flush=True)

        out = {"reference_excel": ref, "live": counts}
        with open(os.path.join(DATA_DIR, "counts_final.json"), "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False, default=str, indent=2)
        with open(os.path.join(DATA_DIR, "cpp_agi_tasks_full.json"), "w", encoding="utf-8") as f:
            json.dump({"rows": all_rows, "counts": counts}, f, ensure_ascii=False, indent=2)
        print("saved counts_final.json + cpp_agi_tasks_full.json", flush=True)

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Live Counts"
            ws.append(["Metric", "Value"])
            for k, v in counts.items():
                if isinstance(v, dict):
                    for kk, vv in v.items():
                        ws.append(["%s [%s]" % (k, kk), vv])
                else:
                    ws.append([k, v])
            ws2 = wb.create_sheet("Live Tasks (E&I&T)")
            cols = ["TaskName", "AssetTag", "Discipline", "TaskState", "ApprovedDate", "ApprovedBy",
                    "System", "Physical Location", "ProcessBreakdown", "TaskType", "TaskModelName", "LoopName"]
            ws2.append(cols)
            write_rows = 0
            for r in all_rows:
                disc = disc_code(r.get("TaskDisciplineSummary") or r.get("TaskDiscipline"))
                if disc not in ("E", "I", "T"):
                    continue
                ws2.append([
                    r.get("TaskName"), r.get("AssetTag"), r.get("TaskDisciplineSummary"),
                    r.get("TaskState"), r.get("ApprovedDate"), r.get("ApprovedBy"),
                    r.get("ProcessBreakdownUp3Identifier"), r.get("PhysicalLocationUp1Summary"),
                    r.get("ProcessBreakdown"), r.get("TaskType"), r.get("TaskModelName"), r.get("LoopName"),
                ])
                write_rows += 1
                if write_rows > 200000:
                    break
            outx = os.path.join(BASE_DIR, "Live_CPP_AGI_EIT_Tasks.xlsx")
            wb.save(outx)
            print("saved:", outx, "rows:", write_rows, flush=True)
        except Exception as e:
            print("excel out err:", str(e)[:200], flush=True)

        ctx.close()


def pos8(s):
    return "PS5" in (s or "")


if __name__ == "__main__":
    sys.exit(main())