"""
PS5 - Apply platform edits into the LATEST 'PS-5 COMPLETIONS DPR SUMMERY' file.

Reads _explorer_state.json (browser edits saved via explorer_server.py) and writes
PUNCH LIST / ITR LIST cell overrides + USER NOTES into the newest
PS-5 COMPLETIONS DPR SUMMERY *.xlsx (whatever its date suffix).

Matching:
  rows   -> by ID column ('Punchlist ID' / 'Task ID')
  columns-> by header tokens (e.g. 'CLOSING DATE' matches 'Workflow - Closing Date')
RFC PROGRESS edits are skipped on purpose (that sheet is formula-driven).

A rolling backup 'PS-5 COMPLETIONS DPR SUMMERY -BACKUP.xlsx' is written first.
"""
import glob
import json
import os
import re
import shutil
import time

import openpyxl
from openpyxl.styles import Font

HOME = os.path.join(os.path.expanduser('~'), 'Downloads')
DL_SUB = os.path.join(HOME,
                      'PS5 - CPP AGI Completion Progress Dashboard_files')
STATE = os.path.join(DL_SUB, '_explorer_state.json')

TARGET_SHEET = {
    'PUNCH LIST': ('DETAILED PUNCH LIST', 'PUNCHLIST ID'),
    'ITR LIST': ('DETAILED ITR LIST', 'TASK ID'),
}


def n(v):
    return '' if v is None else str(v).strip()


def norm(s):
    return re.sub(r'[^A-Z0-9 ]+', ' ', n(s).upper()).split()


def find_latest_summery():
    cands = []
    for folder in (DL_SUB, HOME):
        if not os.path.isdir(folder):
            continue
        for f in glob.glob(os.path.join(folder, '*.xlsx')):
            b = os.path.basename(f)
            if b.startswith('~$') or 'completions dpr summery' not in b.lower():
                continue
            if b.upper().endswith('BACKUP.XLSX'):
                continue
            cands.append((os.path.getmtime(f), f))
    return max(cands)[1] if cands else None


def load_state():
    try:
        with open(STATE, encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}


def build_maps(ws):
    hrow = 1
    raw = {c: n(ws.cell(hrow, c).value) for c in range(1, ws.max_column + 1)}
    idc = next((c for c, v in raw.items()
                if v.upper().startswith(IDH)), 1)
    ids = {}
    for r in range(hrow + 1, ws.max_row + 1):
        rid = n(ws.cell(r, idc).value)
        if rid and rid not in ids:
            ids[rid] = r
    return raw, ids


def match_col(raw, cname):
    want = norm(cname)
    if not want:
        return None
    for c, v in raw.items():
        if norm(v) == want:
            return c
    best = None
    best_score = 0
    for c, v in raw.items():
        have = set(norm(v))
        if have and all(t in have for t in want):
            score = len(want) / max(1, len(have))
            if score > best_score:
                best, best_score = c, score
    return best


def main():
    src = find_latest_summery()
    if not src:
        print('  [summery-sync] no DPR SUMMERY file found - skipped')
        return 0
    st = load_state()
    cells = st.get('cells', {}) or {}
    notes = st.get('notes', {}) or {}
    total_v = sum(len(v) for s, d in cells.items() if s in TARGET_SHEET
                  for v in d.values())
    total_n = sum(len(v) for s, v in notes.items() if s in TARGET_SHEET)
    print(f'  [summery-sync] latest = {os.path.basename(src)}  '
          f'(edits: {total_v} cells / {total_n} notes)')
    if not total_v and not total_n:
        print('  [summery-sync] nothing to apply')
        return 0

    backup = os.path.join(os.path.dirname(src),
                          'PS-5 COMPLETIONS DPR SUMMERY -BACKUP.xlsx')
    try:
        shutil.copy2(src, backup)
    except Exception as e:
        print(f'  [summery-sync] backup failed ({e}) - aborted for safety')
        return 1

    wb = openpyxl.load_workbook(src)
    n_val = n_note = missed = 0
    for psheet, (tgt, idh) in TARGET_SHEET.items():
        global IDH
        IDH = idh
        if tgt not in wb.sheetnames:
            continue
        ws = wb[tgt]
        raw, ids = build_maps(ws)
        note_col = match_col(raw, 'USER NOTES')
        if note_col is None:
            note_col = match_col(raw, 'REMARKS')
        for rid, ed in (cells.get(psheet, {}) or {}).items():
            r = ids.get(n(rid))
            if not r:
                continue
            for cname, val in ed.items():
                cc = match_col(raw, cname)
                if cc is None:
                    missed += 1
                    continue
                ws.cell(r, cc).value = val
                ws.cell(r, cc).font = Font(size=9)
                n_val += 1
        for rid, val in (notes.get(psheet, {}) or {}).items():
            r = ids.get(n(rid))
            if not r or note_col is None:
                continue
            ws.cell(r, note_col).value = val
            ws.cell(r, note_col).font = Font(size=9)
            n_note += 1

    ok = False
    for attempt in range(15):
        try:
            wb.save(src)
            ok = True
            break
        except PermissionError:
            print(f'  [wait] close "{os.path.basename(src)}" in Excel '
                  f'... ({attempt + 1}/15)')
            time.sleep(2)
    wb.close()
    if not ok:
        shutil.copy2(backup, src)
        print('  [summery-sync] file locked - restored backup, nothing changed')
        return 1
    extra = f', {missed} skipped (no matching column)' if missed else ''
    print(f'  [summery-sync] OK -> {n_val} cells, {n_note} notes written'
          f'{extra}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
