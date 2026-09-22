import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
from copy import copy
import os

DIR = os.path.dirname(os.path.abspath(__file__))
DASHBOARD = r"C:\Users\mylap\OneDrive\Desktop\dashboard"

def find_latest(folder, keywords):
    candidates = []
    search_dirs = [folder]
    for extra in [r"C:\Users\mylap\Downloads\asset and punch",
                  r"C:\Users\mylap\Downloads\subsystem and punch"]:
        if os.path.isdir(extra) and extra not in search_dirs:
            search_dirs.append(extra)
    for d in search_dirs:
        if not os.path.isdir(d):
            continue
        for f in os.listdir(d):
            if f.startswith('~$'):
                continue
            low = f.lower()
            if low.endswith('.xlsx') and all(k.lower() in low for k in keywords):
                full = os.path.join(d, f)
                candidates.append((os.path.getmtime(full), full))
    if candidates:
        candidates.sort(key=lambda x: x[0], reverse=True)
        return candidates[0][1]
    return None

FILE = os.path.join(DIR, 'PS5_EIT_Dashboard_CPP_AGI.xlsx')
OV = find_latest(DASHBOARD, ['ovtasks', 'testsplanned'])
INSP = find_latest(DASHBOARD, ['inspection', 'register'])
PUNCH_FILE = find_latest(DASHBOARD, ['eit', 'punch', 'list', 'register']) or os.path.join(DIR, 'PS-5 EIT PUNCH LIST REGISTER.xlsx')
OLD = os.path.join(DIR, 'PS5_EIT_Dashboard_CPP_AGI_old.xlsx')
COMPANY = 'CPP AGI'
DS = {'E - Electrical': 'E', 'Electrical': 'E',
      'I - Instrumentation': 'I', 'I - Instrument': 'I', 'Instrumentation': 'I', 'Instrument': 'I',
      'T - Telecom': 'T', 'Telecom': 'T'}
COLS = ['Task ID', 'Asset - Tag', 'Description', 'Discipline', 'Subsystem', 'Task State',
        'INSTALLATION & PULLING', 'Testing', 'GLANDING & TERMINATION INSPECTION',
        'PUNCH A', 'PUNCH B', 'PUNCH C']

def ss(v):
    if pd.isna(v): return ''
    return str(v).strip()

print('[1/7] Loading ovTasks...')
df = pd.read_excel(OV, sheet_name='Exported from SC')
df['disc'] = df['Discipline (Summary)'].astype(str).str.strip().map(DS).fillna('Other')
eit = df[(df['Responsible Company (Summary)'] == COMPANY) & (df['disc'].isin(['E','I','T']))].copy().reset_index(drop=True)
is_closed = eit['Task State'] == 'Closed'
eit_open = eit[~is_closed]
eit_close = eit[is_closed]
open_tids = set(eit_open['Task ID'].astype(str).str.strip())
close_tids = set(eit_close['Task ID'].astype(str).str.strip())
all_tids = open_tids | close_tids
T_O = len(open_tids)
T_C = len(close_tids)
T_ALL = T_O + T_C
print(f'  {T_ALL} total, {T_O} open, {T_C} closed')

print('[2/7] Task State Overview...')
ts_counts = {}
for state, cnt in eit['Task State'].value_counts().items():
    ts_counts[state] = int(cnt)
states_order = ['To be completed', 'Started (not Completed)', 'Submitted', 'Completed (not Closed)', 'Closed']
status_data = []
for s in states_order:
    cnt = ts_counts.get(s, 0)
    is_close_state = (s == 'Closed')
    o = 0 if is_close_state else cnt
    c = cnt if is_close_state else 0
    pct = f'{c/(o+c)*100:.1f}%' if (o+c) else '0.0%'
    status_data.append((s, o, c, o+c, pct))

print('[3/7] Discipline Breakdown...')
disc_data = []
for d in ['E', 'I', 'T']:
    sub = eit[eit['disc'] == d]
    o = int((sub['Task State'] != 'Closed').sum())
    c = int((sub['Task State'] == 'Closed').sum())
    t = o + c
    pct = f'{c/t*100:.1f}%' if t else '0.0%'
    disc_data.append((d, o, c, t, pct))

print('[4/7] Conformity Check & Static Test...')
cc = eit['Task Type (Name)'].astype(str).str.strip()
is_cc = cc.str.contains('Conformity Check', case=False, na=False)
is_st = cc.str.contains('Static Test', case=False, na=False)
cc_o = int((is_cc & ~is_closed).sum()); cc_c = int((is_cc & is_closed).sum())
cc_t = cc_o + cc_c
other_o = int((~is_cc & ~is_closed).sum()); other_c = int((~is_cc & is_closed).sum())
other_t = other_o + other_c
st_o = int((is_st & ~is_closed).sum()); st_c = int((is_st & is_closed).sum())
st_t = st_o + st_c

print('[5/7] Loading Inspection Register...')
_insp_raw = pd.read_excel(INSP, sheet_name='PS-5 EIT INSPECTION REGISTER', header=None)
_insp_hrow = None
for _i in range(min(200, len(_insp_raw))):
    _v = _insp_raw.iloc[_i, 1]
    if pd.notna(_v) and str(_v).strip() == 'Asset - Tag':
        _insp_hrow = _i
        break
if _insp_hrow is None:
    raise KeyError('Asset - Tag')
insp = pd.read_excel(INSP, sheet_name='PS-5 EIT INSPECTION REGISTER', header=_insp_hrow)
insp['Task ID'] = insp['Task ID'].astype(str).str.strip()
insp['Asset - Tag'] = insp['Asset - Tag'].astype(str).str.strip()
insp_type = insp.iloc[:, 12].astype(str).str.upper()

# Build RFI dicts keyed by TASK ID (primary) and ASSET TAG (fallback)
inst_pull_by_tid, inst_pull_by_tag = {}, {}
testing_by_tid, testing_by_tag = {}, {}
glanding_by_tid, glanding_by_tag = {}, {}
for _, row in insp.iterrows():
    asset_tag = ss(row.get('Asset - Tag', ''))
    tid = ss(row.get('Task ID', ''))
    rfi_i = ss(row.get('QC RFI#', ''))
    it = str(row.iloc[12]).strip().upper() if len(insp.columns) > 12 else ''
    rfi_n = ss(row.iloc[13]) if len(insp.columns) > 13 else ''
    valid_tag = asset_tag not in ('N/A', 'None', '', 'nan')
    valid_tid = tid not in ('nan', '', 'None')
    if it in ('INSTALLATION', 'PULLING') and rfi_i and rfi_i not in ('None', 'nan'):
        if valid_tag:
            inst_pull_by_tag.setdefault(asset_tag, rfi_i)
        if valid_tid:
            inst_pull_by_tid.setdefault(tid, rfi_i)
    if it == 'TESTING' and rfi_i and rfi_i not in ('None', 'nan'):
        if valid_tag:
            testing_by_tag.setdefault(asset_tag, rfi_i)
        if valid_tid:
            testing_by_tid.setdefault(tid, rfi_i)
    if rfi_n and rfi_n not in ('None', 'nan', ''):
        if valid_tag:
            glanding_by_tag.setdefault(asset_tag, rfi_n)
        if valid_tid:
            glanding_by_tid.setdefault(tid, rfi_n)

def rfidict(tid, tag, d_tid, d_tag):
    return d_tid.get(tid) or d_tag.get(tag) or ''

print('[6/7] Loading Punch List...')
punch = pd.read_excel(PUNCH_FILE, sheet_name='MP Register', header=3)
punch_data = {}
for _, row in punch.iterrows():
    vals = list(row.values)
    tag = ss(vals[7]) if len(vals) > 7 else ''
    cat = ss(vals[9]).upper() if len(vals) > 9 else ''
    desc = ss(vals[8]) if len(vals) > 8 else ''
    status = ss(vals[20]).upper() if len(vals) > 20 else ''
    if cat in ('A', 'B', 'C') and status == 'OPEN':
        punch_data.setdefault((tag, cat), []).append(f"{desc} [OPEN]")

def get_punch(tag, cat):
    return ' | '.join(punch_data.get((tag, cat), []))

eit_tags = set(eit['Asset - Tag'].astype(str).str.strip())
punch_counts = {}
for cat in ['A', 'B', 'C']:
    mask = (punch.iloc[:, 9].astype(str).str.strip().str.upper() == cat) & \
           (punch.iloc[:, 7].astype(str).str.strip().isin(eit_tags))
    sub_status = punch.iloc[:, 20].astype(str).str.strip().str.upper()
    o = int((mask & (sub_status == 'OPEN')).sum())
    c = int((mask & (sub_status == 'CLOSED')).sum())
    punch_counts[cat] = (o, c)

# Build OPEN/CLOSE data
out_data = []
for _, row in eit.iterrows():
    tid = ss(row.get('Task ID', ''))
    tag = ss(row.get('Asset - Tag', ''))
    out_data.append({
        'Task ID': tid, 'Asset - Tag': tag,
        'Description': ss(row.get('Asset - Description', '')),
        'Discipline': ss(row.get('disc', '')),
        'Subsystem': ss(row.get('Systemization - Subsystem (Summary)', '')),
        'Task State': ss(row.get('Task State', '')),
    })
out_df = pd.DataFrame(out_data)
out_df['INSTALLATION & PULLING'] = out_df.apply(lambda r: rfidict(r['Task ID'], r['Asset - Tag'], inst_pull_by_tid, inst_pull_by_tag), axis=1)
out_df['Testing'] = out_df.apply(lambda r: rfidict(r['Task ID'], r['Asset - Tag'], testing_by_tid, testing_by_tag), axis=1)
out_df['GLANDING & TERMINATION INSPECTION'] = out_df.apply(lambda r: rfidict(r['Task ID'], r['Asset - Tag'], glanding_by_tid, glanding_by_tag), axis=1)
out_df['PUNCH A'] = out_df['Asset - Tag'].apply(lambda t: get_punch(t, 'A'))
out_df['PUNCH B'] = out_df['Asset - Tag'].apply(lambda t: get_punch(t, 'B'))
out_df['PUNCH C'] = out_df['Asset - Tag'].apply(lambda t: get_punch(t, 'C'))
open_df = out_df[out_df['Task State'] != 'Closed']
close_df = out_df[out_df['Task State'] == 'Closed']

print('  Writing OPEN/CLOSE sheets...')
wb = openpyxl.load_workbook(FILE)
for sn, data_df, title in [
    ('PS-5 EIT TASK OPEN', open_df, 'PS-5 EIT TASK MANAGEMENT - OPEN TASKS'),
    ('PS-5 EIT TASK CLOSE', close_df, 'PS-5 EIT TASK MANAGEMENT - CLOSED TASKS'),
]:
    ws = wb[sn]
    for mg in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mg))
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=12):
        for cell in row:
            cell.value = None
    ws['A1'] = title
    for ci, cn in enumerate(COLS, 1):
        ws.cell(row=2, column=ci, value=cn)
    for ri, (_, row) in enumerate(data_df.iterrows(), 3):
        for ci, cn in enumerate(COLS, 1):
            val = row[cn]
            ws.cell(row=ri, column=ci, value=str(val) if val else '')
    print(f'    {sn}: {len(data_df)} rows')

print('[7/7] Updating Summary...')
ws = wb['Summary']

def pct(a, b):
    return f'{b/(a+b)*100:.1f}%' if (a+b) else '0.0%'

# Section 1: TASK STATE OVERVIEW
for i, (s, o, c, t, p) in enumerate(status_data):
    r = 5 + i
    ws.cell(row=r, column=1, value=s)
    ws.cell(row=r, column=2, value=o)
    ws.cell(row=r, column=3, value=c)
    ws.cell(row=r, column=4, value=t)
    ws.cell(row=r, column=5, value=p)
ws['A10'] = 'TOTAL'; ws['B10'] = T_O; ws['C10'] = T_C; ws['D10'] = T_ALL; ws['E10'] = pct(T_O, T_C)

# Section 2: DISCIPLINE
for i, (d, o, c, t, p) in enumerate(disc_data):
    r = 14 + i
    ws.cell(row=r, column=2, value=o); ws.cell(row=r, column=3, value=c)
    ws.cell(row=r, column=4, value=t); ws.cell(row=r, column=5, value=p)
ws['B17'] = T_O; ws['C17'] = T_C; ws['D17'] = T_ALL; ws['E17'] = pct(T_O, T_C)

# Section 3: CONFORMITY CHECK
ws['B21'] = cc_o; ws['C21'] = cc_c; ws['D21'] = cc_t; ws['E21'] = pct(cc_o, cc_c)
ws['B22'] = other_o; ws['C22'] = other_c; ws['D22'] = other_t; ws['E22'] = pct(other_o, other_c)

# Section 4: STATIC TEST
ws['B26'] = st_o; ws['C26'] = st_c; ws['D26'] = st_t; ws['E26'] = pct(st_o, st_c)

# Section 5: INSTALLATION & PULLING RFI
pull_o = int(open_df['INSTALLATION & PULLING'].ne('').sum()); pull_c = int(close_df['INSTALLATION & PULLING'].ne('').sum())
ws['B30'] = pull_o; ws['C30'] = pull_c; ws['D30'] = pull_o + pull_c; ws['E30'] = pct(pull_o, pull_c)
ws['B31'] = T_O - pull_o; ws['C31'] = T_C - pull_c; ws['D31'] = T_ALL - pull_o - pull_c; ws['E31'] = pct(T_O - pull_o, T_C - pull_c)
ws['B32'] = T_O; ws['C32'] = T_C; ws['D32'] = T_ALL; ws['E32'] = pct(T_O, T_C)

# Section 6: TESTING RFI
test_o = int(open_df['Testing'].ne('').sum()); test_c = int(close_df['Testing'].ne('').sum())
ws['B36'] = test_o; ws['C36'] = test_c; ws['D36'] = test_o + test_c; ws['E36'] = pct(test_o, test_c)
ws['B37'] = T_O - test_o; ws['C37'] = T_C - test_c; ws['D37'] = T_ALL - test_o - test_c; ws['E37'] = pct(T_O - test_o, T_C - test_c)
ws['B38'] = T_O; ws['C38'] = T_C; ws['D38'] = T_ALL; ws['E38'] = pct(T_O, T_C)

# Section 7: GLANDING
glan_o = int(open_df['GLANDING & TERMINATION INSPECTION'].ne('').sum()); glan_c = int(close_df['GLANDING & TERMINATION INSPECTION'].ne('').sum())
ws['B42'] = glan_o; ws['C42'] = glan_c; ws['D42'] = glan_o + glan_c; ws['E42'] = pct(glan_o, glan_c)
ws['B43'] = T_O - glan_o; ws['C43'] = T_C - glan_c; ws['D43'] = T_ALL - glan_o - glan_c; ws['E43'] = pct(T_O - glan_o, T_C - glan_c)
ws['B44'] = T_O; ws['C44'] = T_C; ws['D44'] = T_ALL; ws['E44'] = pct(T_O, T_C)

# Section 8: PUNCH
po_a, pc_a = punch_counts['A']; po_b, pc_b = punch_counts['B']; po_c, pc_c = punch_counts['C']
ws['B48'] = po_a; ws['C48'] = pc_a; ws['D48'] = po_a + pc_a; ws['E48'] = pct(po_a, pc_a)
ws['B49'] = po_b; ws['C49'] = pc_b; ws['D49'] = po_b + pc_b; ws['E49'] = pct(po_b, pc_b)
ws['B50'] = po_c; ws['C50'] = pc_c; ws['D50'] = po_c + pc_c; ws['E50'] = pct(po_c, pc_c)
pt_o = po_a + po_b + po_c; pt_c = pc_a + pc_b + pc_c
ws['B51'] = pt_o; ws['C51'] = pt_c; ws['D51'] = pt_o + pt_c; ws['E51'] = pct(pt_o, pt_c)

print('  Summary updated')
print('  Copying formatting...')
if os.path.exists(OLD):
    old_wb = openpyxl.load_workbook(OLD)
    for sn in old_wb.sheetnames:
        old_ws = old_wb[sn]
        new_ws = wb[sn]
        for mg in list(new_ws.merged_cells.ranges):
            new_ws.unmerge_cells(str(mg))
        for mg in old_ws.merged_cells.ranges:
            new_ws.merge_cells(str(mg))
        for cl in 'ABCDEFGHIJKLMNOPQRSTUVWXYZ':
            od = old_ws.column_dimensions.get(cl)
            if od and od.width:
                new_ws.column_dimensions[cl].width = od.width
        for r in range(1, old_ws.max_row + 1):
            od = old_ws.row_dimensions.get(r)
            if od and od.height:
                new_ws.row_dimensions[r].height = od.height
        for r in range(1, 3):
            for c in range(1, 13):
                oc = old_ws.cell(row=r, column=c)
                nc = new_ws.cell(row=r, column=c)
                if oc.font: nc.font = copy(oc.font)
                if oc.fill: nc.fill = copy(oc.fill)
                if oc.alignment: nc.alignment = copy(oc.alignment)
                if oc.border: nc.border = copy(oc.border)
        if old_ws.sheet_properties.tabColor:
            new_ws.sheet_properties.tabColor = copy(old_ws.sheet_properties.tabColor)
        if sn != 'Summary' and old_ws.max_row >= 4:
            df_ = [copy(old_ws.cell(row=3, column=c).font) for c in range(1, 13)]
            dfl = [copy(old_ws.cell(row=3, column=c).fill) for c in range(1, 13)]
            dla = [copy(old_ws.cell(row=3, column=c).alignment) for c in range(1, 13)]
            dlb = [copy(old_ws.cell(row=3, column=c).border) for c in range(1, 13)]
            af_ = [copy(old_ws.cell(row=4, column=c).font) for c in range(1, 13)]
            afl = [copy(old_ws.cell(row=4, column=c).fill) for c in range(1, 13)]
            afla = [copy(old_ws.cell(row=4, column=c).alignment) for c in range(1, 13)]
            aflb = [copy(old_ws.cell(row=4, column=c).border) for c in range(1, 13)]
            for r in range(3, new_ws.max_row + 1):
                idx = (r - 3) % 2
                for c in range(1, 13):
                    cell = new_ws.cell(row=r, column=c)
                    src = [df_, af_][idx]; sf = [dfl, afl][idx]
                    sa = [dla, afla][idx]; sb = [dlb, aflb][idx]
                    cell.font = src[c-1]; cell.fill = sf[c-1]
                    cell.alignment = sa[c-1]; cell.border = sb[c-1]
    print('  Formatting done')
else:
    print('  Old file not found, skip formatting')

wb.save(FILE)
print(f'\nDone! {T_ALL} total, {T_O} open, {T_C} closed')
input('Press Enter to close...')
