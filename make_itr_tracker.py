import pandas as pd
import os, shutil, tempfile

BASE = r"C:\Users\mylap\OneDrive\Desktop\dashboard"
TMP = tempfile.gettempdir()

def find_file(patterns):
    for f in os.listdir(BASE):
        low = f.lower()
        if all(p.lower() in low for p in patterns):
            return os.path.join(BASE, f)
    return None

def copy_to_tmp(path):
    """Copy file to temp to bypass Excel lock."""
    dst = os.path.join(TMP, os.path.basename(path))
    try:
        shutil.copy2(path, dst)
    except PermissionError:
        return None
    return dst

# ---- Find files ----
ov_path = find_file(['ovtasks'])
punch_path = find_file(['eit', 'punch', 'list', 'register']) or find_file(['punch', 'list', 'register'])
rfi_path = find_file(['inspection', 'register'])

if not ov_path:
    print("ERROR: ovTasks not found"); exit(1)

print(f"ovTasks       : {ov_path}")
print(f"Punch List    : {punch_path or 'NOT FOUND'}")
print(f"Inspection Reg: {rfi_path or 'NOT FOUND'}")

# ---- Copy to temp ----
print("\n[0] Copying files to temp (bypass Excel lock)...")
ov_tmp = copy_to_tmp(ov_path)
if punch_path:
    punch_tmp = copy_to_tmp(punch_path)
else:
    punch_tmp = None
if rfi_path:
    rfi_tmp = copy_to_tmp(rfi_path)
else:
    rfi_tmp = None

if not ov_tmp:
    print("ERROR: Cannot copy ovTasks"); exit(1)

# ============================================================
# 1) Read Construction (ovTasks)
# ============================================================
print("\n[1] Reading Construction tasks (ovTasks)...")
cols = ['Task ID', 'Asset - Tag', 'Description', 'Discipline (Summary)',
        'Responsible Company (Summary)', 'Systemization - Subsystem (Summary)',
        'Closing Date', 'Task State']
ov = pd.read_excel(ov_tmp, sheet_name='Exported from SC', usecols=cols)
print(f"  Total: {len(ov)}")

cpp = ov[ov['Responsible Company (Summary)'] == 'CPP AGI'].copy()
print(f"  CPP AGI: {len(cpp)}")

disc_map = {
    'E - Electrical': 'E', 'E - Elect': 'E', 'Electrical': 'E', 'E': 'E',
    'I - Instrumentation': 'I', 'I - Instrument': 'I', 'Instrumentation': 'I', 'Instrument': 'I', 'I': 'I',
    'T - Telecom': 'T', 'Telecom': 'T', 'T': 'T',
}
cpp['disc'] = cpp['Discipline (Summary)'].astype(str).str.strip().map(disc_map).fillna('Other')
cpp = cpp[cpp['disc'].isin(['E', 'I', 'T'])].copy()
print(f"  E/I/T: {len(cpp)}")

# ============================================================
# 2) Inspection Register
# ============================================================
rfi_data = {}
if rfi_tmp:
    print("\n[2] Reading Inspection Register...")
    ir = pd.read_excel(rfi_tmp, sheet_name='PS-5 EIT INSPECTION REGISTER', header=5)
    ir = ir.dropna(subset=['Asset - Tag'])
    ir['disc'] = ir['Discipline'].astype(str).str.strip().map(disc_map).fillna('Other')
    ir = ir[ir['disc'].isin(['E', 'I', 'T'])].copy()
    print(f"  Rows (E/I/T): {len(ir)}")

    def norm_status(v):
        if not isinstance(v, str): return 'No RFI Yet'
        v = v.strip().lower()
        if 'punch' in v: return 'Accepted with Punch'
        if 'accept' in v: return 'Accepted'
        if 'hold' in v: return 'Hold by EACOP'
        if 'open' in v: return 'Open'
        return 'Other'
    ir['rfi_status'] = ir['STATUS OF RFI'].apply(norm_status)
    ir['has_rfi'] = ir['QC RFI#'].notna() | ir['QC RFI#.1'].notna()
    for _, row in ir.iterrows():
        tag = str(row['Asset - Tag']).strip()
        if tag not in rfi_data:
            rfi_data[tag] = {'has_rfi': row['has_rfi'], 'rfi_status': row['rfi_status']}
    print(f"  Unique tags: {len(rfi_data)}")

# ============================================================
# 3) Punch List
# ============================================================
punch_data = {}
if punch_tmp:
    print("\n[3] Reading Punch List...")
    pl = pd.read_excel(punch_tmp, sheet_name='MP Register', header=None, skiprows=6)
    for _, row in pl.iterrows():
        if pd.isna(row.iloc[1]): continue
        asset = str(row.iloc[7]).strip() if pd.notna(row.iloc[7]) else ''
        if not asset or asset in ('N/A', '-', ''): continue
        cat = str(row.iloc[9]).strip().upper() if pd.notna(row.iloc[9]) else ''
        st = str(row.iloc[20]).strip() if pd.notna(row.iloc[20]) else ''
        st = st.upper()
        if asset:
            if asset not in punch_data:
                punch_data[asset] = {'A': [], 'B': [], 'C': []}
            if cat in ('A', 'B', 'C'):
                punch_data[asset][cat].append(st)
    print(f"  Assets with punch: {len(punch_data)}")

def punch_summary(cats):
    parts = []
    for c in ['A', 'B', 'C']:
        items = cats.get(c, [])
        if not items: continue
        closed = sum(1 for s in items if s == 'CLOSED')
        parts.append(f"{c}:{closed}/{len(items)}")
    return ', '.join(parts) if parts else '-'

def has_open_punch(cats):
    for c in ['A', 'B', 'C']:
        items = cats.get(c, [])
        if items and any(s != 'CLOSED' for s in items):
            return True
    return False

# ============================================================
# 4) Build Tracker
# ============================================================
print("\n[4] Building ITR Tracker...")
rows = []
for _, row in cpp.iterrows():
    tag = str(row.get('Asset - Tag', '')).strip() if pd.notna(row.get('Asset - Tag')) else ''
    if not tag:
        continue
    tid = str(row.get('Task ID', ''))
    desc = str(row.get('Description', ''))[:120] if pd.notna(row.get('Description')) else ''
    tstate = str(row.get('Task State', ''))
    disc = row['disc']
    subsystem = str(row.get('Systemization - Subsystem (Summary)', ''))[:60] if pd.notna(row.get('Systemization - Subsystem (Summary)', '')) else ''

    insp = rfi_data.get(tag, {})
    has_rfi = insp.get('has_rfi', False)
    rfi_status = insp.get('rfi_status', 'No RFI Yet')

    punch = punch_data.get(tag, {})
    punch_sum = punch_summary(punch)
    open_punch = has_open_punch(punch)

    # Variation
    if tstate == 'Closed' and has_rfi and rfi_status in ('Accepted', 'Accepted with Punch') and not open_punch:
        var = 'Complete'
    elif tstate == 'Closed' and not has_rfi:
        var = 'Closed but No RFI'
    elif tstate == 'Closed' and has_rfi:
        var = f'Closed + {rfi_status}'
    elif tstate != 'Closed' and has_rfi:
        var = 'Open + RFI Submitted'
    elif tstate == 'Submitted' and not has_rfi:
        var = 'Submitted No RFI'
    elif not has_rfi:
        var = 'Not Started'
    else:
        var = 'Other'

    rows.append({
        'Asset Tag': tag,
        'Task ID': tid,
        'Discipline': disc,
        'Subsystem': subsystem,
        'Description': desc,
        'Construction': tstate,
        'RFI?': 'Yes' if has_rfi else 'No',
        'RFI Status': rfi_status,
        'Punch': punch_sum,
        'Open Punch': 'Yes' if open_punch else '',
        'Variation': var,
    })

tracker = pd.DataFrame(rows)
print(f"  Rows: {len(tracker)}")

# ============================================================
# 5) Summary
# ============================================================
print("\n[5] Variation Summary:")
for v, c in tracker['Variation'].value_counts().items():
    print(f"  {v}: {c}")

# ============================================================
# 6) Write Excel
# ============================================================
out = os.path.join(BASE, 'CPP_AGI_ITR_Tracker.xlsx')
print(f"\n[6] Writing {out}...")

from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

with pd.ExcelWriter(out, engine='openpyxl') as writer:
    # --- Sheet 1: Detail ---
    tracker.to_excel(writer, sheet_name='ITR Tracker', index=False)
    ws = writer.sheets['ITR Tracker']
    for col in ws.columns:
        letter = col[0].column_letter
        mx = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[letter].width = min(mx + 2, 45)
    ws.auto_filter.ref = ws.dimensions

    # Color rows
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
        var = str(r[10].value or '')  # Variation col
        if var == 'Complete':
            fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        elif 'Closed but No RFI' in var or 'Submitted No RFI' in var:
            fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
        elif 'Open + RFI' in var:
            fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
        elif var == 'Not Started':
            fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
        elif 'Closed +' in var:
            fill = PatternFill(start_color='FFD966', end_color='FFD966', fill_type='solid')
        else:
            fill = None
        if fill:
            for cell in r:
                cell.fill = fill

    # --- Sheet 2: Variation Summary ---
    var_sum = tracker.groupby('Variation').agg(Count=('Asset Tag', 'count')).reset_index()
    var_sum.to_excel(writer, sheet_name='By Variation', index=False)
    ws2 = writer.sheets['By Variation']
    for col in ws2.columns:
        letter = col[0].column_letter
        mx = max(len(str(c.value or '')) for c in col)
        ws2.column_dimensions[letter].width = mx + 2

    # --- Sheet 3: Discipline Summary ---
    disc_var = tracker.groupby(['Discipline', 'Variation']).agg(Count=('Asset Tag', 'count')).reset_index()
    dp = disc_var.pivot(index='Discipline', columns='Variation', values='Count').fillna(0).astype(int)
    dp['Total'] = dp.sum(axis=1)
    dp.to_excel(writer, sheet_name='By Discipline')
    ws3 = writer.sheets['By Discipline']

    # --- Sheet 4: Subsystem Summary ---
    sub_var = tracker.groupby(['Subsystem', 'Variation']).agg(Count=('Asset Tag', 'count')).reset_index()
    sp = sub_var.pivot(index='Subsystem', columns='Variation', values='Count').fillna(0).astype(int)
    sp['Total'] = sp.sum(axis=1)
    sp = sp.sort_values('Total', ascending=False)
    sp.to_excel(writer, sheet_name='By Subsystem')
    ws4 = writer.sheets['By Subsystem']

    # --- Sheet 5: Construction vs RFI cross-tab ---
    cross = tracker.groupby(['Construction', 'RFI?']).agg(Count=('Asset Tag', 'count')).reset_index()
    cp = cross.pivot(index='Construction', columns='RFI?', values='Count').fillna(0).astype(int)
    cp['Total'] = cp.sum(axis=1)
    cp.to_excel(writer, sheet_name='Construction vs RFI')

print(f"\nDone! File: {out}")
print(f"  Total rows: {len(tracker)}")
print(f"  Sheets: ITR Tracker, By Variation, By Discipline, By Subsystem, Construction vs RFI")
